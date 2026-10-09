# -*- coding: utf-8 -*-
"""Fit and style Chang PPT tables; preserve fixed pages and configured layouts."""
import argparse
import json
import os
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SLACK, PAD = 1.04, 0.08
HEADER_FILL, HEADER_TEXT, BORDER_HEX = '1B2A4A', 'FFFFFF', '7F8FA9'
ALLOWED = {'1B2A4A','2563EB','0F766E','B45309','355C7D','7C3AED','3A86FF'}
_FONTS = {}

class TableFitError(ValueError):
    pass

def find_font(bold=False):
    roots = [os.environ.get('CHANG_PPT_FONT_DIR',''),
             os.path.join(os.environ.get('LOCALAPPDATA',''), 'Microsoft','Windows','Fonts'),
             r'C:\Windows\Fonts', str(Path.home()/'Library/Fonts'),
             str(Path.home()/'.local/share/fonts'), '/usr/share/fonts']
    stem = 'Pretendard-Bold' if bold else 'Pretendard-Regular'
    for root in roots:
        if not root or not Path(root).is_dir():
            continue
        for ext in ('.ttf','.otf'):
            hits = sorted(Path(root).rglob(stem+ext))
            if hits:
                return str(hits[0])
    raise FileNotFoundError(stem+' font missing; set CHANG_PPT_FONT_DIR')

def font(size_pt, bold):
    key = (round(size_pt,1),bold)
    if key not in _FONTS:
        _FONTS[key] = ImageFont.truetype(find_font(bold),round(size_pt*10))
    return _FONTS[key]

def cell_max_line(cell):
    best = 0.0
    for para in cell.text_frame.paragraphs:
        width = 0.0
        for run in para.runs:
            size = run.font.size.pt if run.font.size else 14.0
            bold = bool(run.font.bold if run.font.bold is not None else para.font.bold)
            parts = run.text.replace('\v','\n').split('\n')
            for i, part in enumerate(parts):
                width += font(size,bold).getlength(part)/10/72
                if i < len(parts)-1:
                    best = max(best,width)
                    width = 0.0
        best = max(best,width)
    return best*SLACK

def set_fill(pr, color):
    for tag in ('a:solidFill','a:noFill','a:gradFill','a:pattFill','a:blipFill','a:grpFill'):
        for child in pr.findall(qn(tag)):
            pr.remove(child)
    sf=etree.SubElement(pr,qn('a:solidFill'))
    etree.SubElement(sf,qn('a:srgbClr')).set('val',color)

def set_borders(pr):
    for tag in ('a:lnL','a:lnR','a:lnT','a:lnB'):
        for child in pr.findall(qn(tag)):
            pr.remove(child)
    for i,tag in enumerate(('a:lnL','a:lnR','a:lnT','a:lnB')):
        ln=etree.Element(qn(tag),w=str(int(0.5*12700)),cap='flat',cmpd='sng',algn='ctr')
        set_fill(ln,BORDER_HEX)
        etree.SubElement(ln,qn('a:prstDash')).set('val','solid')
        pr.insert(i,ln)

def set_typeface(pr):
    for tag in ('a:latin','a:ea','a:cs'):
        child=pr.find(qn(tag))
        if child is None:
            child=etree.SubElement(pr,qn(tag))
        child.set('typeface','Pretendard')

def style_cells(shape, config):
    header=config.get('header_fill',HEADER_FILL).lstrip('#').upper()
    if header not in ALLOWED:
        raise ValueError('Unapproved header color: '+header)
    preserve=config.get('preserve_cell_fills',False)
    for ri,row in enumerate(shape.table.rows):
        for cell in row.cells:
            if cell.is_spanned or cell.is_merge_origin:
                raise ValueError('Merged cells require manual layout; not supported by automatic fitter')
            cell.margin_left=cell.margin_right=Inches(0.1)
            cell.margin_top=cell.margin_bottom=Inches(0.03)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            pr=cell._tc.get_or_add_tcPr()
            if not preserve or pr.find(qn('a:solidFill')) is None:
                set_fill(pr,header if ri==0 else ('F5F5F7' if ri%2 else 'FFFFFF'))
            set_borders(pr)
            for para in cell.text_frame.paragraphs:
                para.font.name='Pretendard'
                para.font.size=Pt(14)
                if ri==0:
                    para.font.bold=True
                set_typeface(para._p.get_or_add_pPr().get_or_add_defRPr())
                for run in para.runs:
                    run.font.name='Pretendard'
                    run.font.size=Pt(14)
                    if ri==0:
                        run.font.bold=True
                    if not config.get('preserve_text_colors',False):
                        run.font.color.rgb=RGBColor.from_string(HEADER_TEXT if ri==0 else '2C3E5A')
                    rp=run._r.get_or_add_rPr()
                    set_typeface(rp)
                    if run.text and run.text.isascii():
                        rp.set('lang','en-US')
                end=para._p.find(qn('a:endParaRPr'))
                if end is None:
                    end=etree.SubElement(para._p,qn('a:endParaRPr'))
                set_typeface(end)
                end.set('sz','1400')

def fit_columns(shape, slide_width, config):
    left=float(config.get('area_left',0.65))
    width=float(config.get('area_width',slide_width/914400-left-0.50))
    if left < 0 or width<=0 or left+width>slide_width/914400+1e-5:
        raise ValueError('Invalid table area')
    protect=set(config.get('protected_columns',[]))
    if any(not isinstance(i,int) or i<0 or i>=len(shape.table.columns) for i in protect):
        raise ValueError('Invalid protected column index')
    need=[max(cell_max_line(row.cells[i]) for row in shape.table.rows)+0.2+PAD
          for i in range(len(shape.table.columns))]
    floors=[w if i in protect else min(w,max(0.6,w*0.5)) for i,w in enumerate(need)]
    if sum(floors)>width+1e-6:
        raise TableFitError('Table cannot fit without violating protected widths; shorten or split content')
    while sum(need)>width+1e-6:
        candidates=[i for i,w in enumerate(need) if w>floors[i]+1e-6]
        if not candidates:
            raise TableFitError('Shorten or split table')
        i=max(candidates,key=lambda j:need[j])
        need[i]-=min(sum(need)-width,need[i]-floors[i])
    mode=config.get('align','center')
    if mode not in ('center','preserve'):
        raise ValueError('align must be center or preserve')
    old_left=shape.left
    total=Inches(sum(need))
    new_left=int(Inches(left)+(Inches(width)-total)/2) if mode=='center' else old_left
    if new_left<Inches(left)-1 or new_left+total>Inches(left+width)+1:
        raise TableFitError('Preserved table position exceeds configured area')
    for col,w in zip(shape.table.columns,need):
        col.width=Inches(w)
    shape.width=total
    shape.left=new_left
    return need,new_left-old_left

def apply_table(shape, slide, slide_width, config):
    # Configuration identifies companions explicitly; never guess from nearby shapes.
    names=config.get('companions',[])
    companions=[]
    for name in names:
        matches=[s for s in slide.shapes if s.name==name]
        if len(matches)!=1 or matches[0] is shape:
            raise ValueError('Companion name must identify one other shape: '+name)
        companions.append(matches[0])
    style_cells(shape,config)
    widths,dx=fit_columns(shape,slide_width,config)
    for companion in companions:
        companion.left+=dx
    return widths

def main(src,dst,config=None):
    prs=Presentation(src)
    settings=config or {}
    tables=settings.get('tables',{})
    fixed={1,2,len(prs.slides)}
    seen=set()
    for num,slide in enumerate(prs.slides,1):
        if num in fixed:
            continue
        for index,shape in enumerate([s for s in slide.shapes if s.has_table],1):
            key=f'{num}:{index}'
            seen.add(key)
            item=tables.get(key,{})
            if item.get('skip',False):
                continue
            widths=apply_table(shape,slide,prs.slide_width,item)
            print(f'slide {num} table {index}: {sum(widths):.2f} in; check row height and rendering')
    unused=set(tables)-seen
    if unused:
        raise ValueError('Unknown or fixed-page table settings: '+', '.join(sorted(unused)))
    prs.save(dst)
    print('saved:',dst)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('src')
    parser.add_argument('dst')
    parser.add_argument('--config',help='JSON configuration; see references/table-contract.md')
    args=parser.parse_args()
    try:
        cfg=json.loads(Path(args.config).read_text(encoding='utf-8')) if args.config else {}
        main(args.src,args.dst,cfg)
    except (ValueError,FileNotFoundError) as error:
        parser.exit(2,str(error)+'\n')
