# -*- coding: utf-8 -*-
"""chang-presentation-brand — 표 규격 자동 적용.

  python brand_tables.py <입력.pptx> <출력.pptx>

규칙 (SKILL.md「표 규격」):
  1. 열 폭 = 셀 안 가장 긴 한 줄의 Pretendard 실측 폭 × 1.04 + 셀 여백 + 0.08"
  2. 표 폭 = 열 폭 합, 슬라이드 가로 정중앙
  3. 합이 12.33" 초과 시 가장 넓은 열부터 축소 (하한: 자연 폭 50%, 최소 2.0")
  4. 헤더 행 #1B2A4A + 흰 글자, 모든 셀 0.5pt #7F8FA9 실선 테두리
  5. 순수 ASCII 런에 lang="en-US"
"""
import os
import sys

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

SLIDE_W = 13.3333
AREA_L, AREA_R = 0.50, 12.83
MAX_W = AREA_R - AREA_L            # 12.33"
SLACK = 1.04                       # PIL 측정 → PowerPoint 실측 보정
PAD = 0.08                         # 고정 여유(inch)
HEADER_FILL = "1B2A4A"
HEADER_TEXT = "FFFFFF"
BORDER_HEX = "7F8FA9"
BORDER_EMU = int(0.5 * 12700)      # 0.5pt
SCALE = 10                         # 10배 크기로 렌더해 소수점 정밀도 확보

FONT_DIRS = [
    os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Windows", "Fonts"),
    r"C:\Windows\Fonts",
    os.path.expanduser("~/Library/Fonts"),
    "/usr/share/fonts",
]


def find_font(name: str) -> str:
    for d in FONT_DIRS:
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(f"{name} 을(를) 찾을 수 없습니다. Pretendard를 설치하세요.")


_FONTS: dict[tuple[float, bool], ImageFont.FreeTypeFont] = {}


def font(size_pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    key = (round(size_pt, 1), bold)
    if key not in _FONTS:
        name = "Pretendard-Bold.ttf" if bold else "Pretendard-Regular.ttf"
        _FONTS[key] = ImageFont.truetype(find_font(name), int(size_pt * SCALE))
    return _FONTS[key]


def line_width_in(text: str, size_pt: float, bold: bool) -> float:
    if not text:
        return 0.0
    return font(size_pt, bold).getlength(text) / SCALE / 72.0


def cell_insets_in(cell) -> float:
    pr = cell._tc.find(qn("a:tcPr"))
    l = int(pr.get("marL", 91440)) if pr is not None else 91440
    r = int(pr.get("marR", 91440)) if pr is not None else 91440
    return (l + r) / 914400.0


def cell_max_line(cell) -> float:
    best = 0.0
    for para in cell.text_frame.paragraphs:
        size, bold = 14.0, False
        for r in para.runs:
            if r.font.size:
                size = max(size, r.font.size.pt)
            bold = bold or bool(r.font.bold)
        text = "".join(r.text for r in para.runs)
        for line in text.replace("\v", "\n").split("\n"):
            best = max(best, line_width_in(line.strip(), size, bold) * SLACK)
    return best


def set_latin_lang(cell) -> None:
    for r in cell._tc.iter(qn("a:r")):
        t = r.find(qn("a:t"))
        if t is None or not t.text or not t.text.isascii():
            continue
        rPr = r.find(qn("a:rPr"))
        if rPr is None:
            rPr = r.makeelement(qn("a:rPr"), {})
            r.insert(0, rPr)
        rPr.set("lang", "en-US")


def set_fill(tcPr, hex_: str) -> None:
    for e in tcPr.findall(qn("a:solidFill")) + tcPr.findall(qn("a:noFill")):
        tcPr.remove(e)
    sf = etree.SubElement(tcPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", hex_)


def set_borders(tcPr) -> None:
    tags = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")
    for tag in tags:
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    fill = tcPr.find(qn("a:solidFill"))       # 스키마 순서: ln* 이 fill 앞
    idx = list(tcPr).index(fill) if fill is not None else len(tcPr)
    for i, tag in enumerate(tags):
        ln = etree.Element(qn(tag), w=str(BORDER_EMU), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", BORDER_HEX)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(idx + i, ln)


def fit_columns(shape) -> tuple[list[float], str]:
    tbl = shape.table
    need = []
    for ci, _ in enumerate(tbl.columns):
        w = 0.0
        for row in tbl.rows:
            cell = row.cells[ci]
            set_latin_lang(cell)
            w = max(w, cell_max_line(cell) + cell_insets_in(cell))
        need.append(w + PAD)
    note = "fit"
    if sum(need) > MAX_W:
        floor = [max(2.0, w * 0.5) for w in need]
        while sum(need) > MAX_W + 1e-6:
            i = max(range(len(need)), key=lambda j: need[j])
            cut = min(sum(need) - MAX_W, need[i] - floor[i])
            if cut <= 0:
                k = MAX_W / sum(need)
                need = [w * k for w in need]
                break
            need[i] -= cut
        note = "capped"
    for col, w in zip(tbl.columns, need):
        col.width = Emu(int(Inches(w)))
    shape.width = Emu(int(Inches(sum(need))))
    shape.left = Emu(int((Inches(SLIDE_W) - shape.width) / 2))
    return need, note


def style_cells(shape) -> None:
    for ri, row in enumerate(shape.table.rows):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            if ri == 0:
                set_fill(tcPr, HEADER_FILL)
                for para in cell.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.color.rgb = RGBColor.from_string(HEADER_TEXT)
            set_borders(tcPr)


def main(src: str, dst: str) -> None:
    prs = Presentation(src)
    for num, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            need, note = fit_columns(shape)
            style_cells(shape)
            cols = " ".join(f"{w:.2f}" for w in need)
            print(f"slide {num:2d} {shape.name:<10} {sum(need):5.2f}\"  cols={cols}  [{note}]")
    prs.save(dst)
    print("saved:", dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
