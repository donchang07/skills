"""
표지 · 내소개는 v1 원본을 복원, 콘텐츠 가이드만 v2.0 패턴 유지.
모든 슬라이드 배경을 white로, 우상단 삼각형은 추가 90° 시계방향 회전.
"""

import os
import sys
import shutil

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
V1_BACKUP = os.path.join(HERE, '장동인 PPT템플릿_v1_backup.pptx')
OUTPUT = os.path.join(HERE, 'template.pptx')

# v2.0 색 (cream → white로 교체)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HERITAGE_NAVY = RGBColor(0x1B, 0x2A, 0x4A)
HERITAGE_BLUE = RGBColor(0x3A, 0x86, 0xFF)
DIAGONAL_LIGHT = RGBColor(0xD8, 0xE2, 0xF0)
INK = RGBColor(0x0A, 0x0A, 0x0A)
BODY = RGBColor(0x3A, 0x3A, 0x3A)
MUTED = RGBColor(0x6A, 0x6A, 0x6A)
ON_PRIMARY = RGBColor(0xFF, 0xFF, 0xFF)
BRAND_TEAL = RGBColor(0x1A, 0x3A, 0x3A)
BRAND_PINK = RGBColor(0xFF, 0x4D, 0x8B)
HAIRLINE = RGBColor(0xE5, 0xE5, 0xE5)

SLIDE_W, SLIDE_H = 12192000, 6858000


def remove_all_shapes(slide):
    sp_tree = slide.shapes._spTree
    targets = (
        sp_tree.findall(qn('p:sp'))
        + sp_tree.findall(qn('p:pic'))
        + sp_tree.findall(qn('p:cxnSp'))
        + sp_tree.findall(qn('p:graphicFrame'))
    )
    for sp in targets:
        sp_tree.remove(sp)


def add_rect(slide, x, y, w, h, rgb):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb
    shp.line.fill.background()
    return shp


def add_text(slide, text, x, y, w, h, *, font='Pretendard Variable',
             size=14, bold=False, rgb=BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb
    return tb


def add_multiline(slide, lines, x, y, w, h, *, size=14, rgb=BODY):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.55
        run = p.add_run()
        run.text = line
        run.font.name = 'Pretendard Variable'
        run.font.size = Pt(size)
        run.font.color.rgb = rgb
    return tb


def build_guide_slide(slide):
    """슬라이드 3: 콘텐츠 가이드 — v2.0 비교표 패턴 (white 배경, 삼각형 +90°)."""
    remove_all_shapes(slide)

    # Heritage frame
    add_rect(slide, 0, 0, 360000, SLIDE_H, HERITAGE_NAVY)
    add_rect(slide, 486000, 0, 46800, 1051200, HERITAGE_NAVY)
    add_rect(slide, 360000, 756000, 12240000, 46800, HERITAGE_NAVY)

    # 우상단 대각선 삼각형 — 기존 90° + 추가 90° clockwise = 180°
    diag = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE,
                                  11530000, 0, 660000, 660000)
    diag.rotation = 180
    diag.fill.solid()
    diag.fill.fore_color.rgb = DIAGONAL_LIGHT
    diag.line.fill.background()

    # 하단 라인
    add_rect(slide, 360000, 6829800, 11832000, 28200, HERITAGE_BLUE)

    # 제목 / 부제목
    add_text(slide, '콘텐츠 가이드',
             360000, 270000, 11472000, 432000,
             size=20, bold=True, rgb=INK, align=PP_ALIGN.CENTER)
    add_text(slide, '슬라이드 타이포그래피 · 다이어그램 표준',
             864000, 882000, 10464000, 378000,
             size=16, bold=True, rgb=HERITAGE_BLUE, align=PP_ALIGN.CENTER)

    # 좌측 STANDARD 카드 (white + hairline)
    L_x, L_y, L_w, L_h = 1152000, 2016000, 4392000, 3960000
    card_l = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, L_x, L_y, L_w, L_h)
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = WHITE
    card_l.line.color.rgb = HAIRLINE
    card_l.line.width = Pt(0.75)
    add_text(slide, 'STANDARD', L_x + 360000, L_y + 360000, L_w - 720000, 270000,
             font='Inter', size=10, bold=True, rgb=MUTED)
    add_text(slide, '다이어그램 (보통)',
             L_x + 360000, L_y + 666000, L_w - 720000, 540000,
             size=24, bold=True, rgb=INK)
    add_rect(slide, L_x + 360000, L_y + 1278000, 360000, 18000, HAIRLINE)
    add_multiline(slide, [
        '• Headline — Pretendard 20pt, Bold, ink',
        '• 소제목   — Pretendard 16pt, Bold, navy',
        '• 본문     — Pretendard 14pt, Regular, body',
        '• 카드 배경 — white',
        '• 테두리   — 1px hairline (#e5e5e5)',
    ], L_x + 360000, L_y + 1404000, L_w - 720000, L_h - 1764000,
       size=13, rgb=BODY)

    # 우측 EMPHASIS 카드 (teal + pink Featured 배지)
    R_x, R_y, R_w, R_h = 6660000, 2016000, 4392000, 3960000
    add_rect(slide, R_x, R_y, R_w, R_h, BRAND_TEAL)

    badge_x = R_x + R_w - 1296000
    badge_y = R_y - 162000
    add_rect(slide, badge_x, badge_y, 1188000, 324000, BRAND_PINK)
    add_text(slide, 'FEATURED', badge_x, badge_y, 1188000, 324000,
             font='Inter', size=10, bold=True, rgb=ON_PRIMARY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    add_text(slide, 'EMPHASIS', R_x + 360000, R_y + 360000, R_w - 720000, 270000,
             font='Inter', size=10, bold=True, rgb=BRAND_PINK)
    add_text(slide, '다이어그램 (강조)',
             R_x + 360000, R_y + 666000, R_w - 720000, 540000,
             size=24, bold=True, rgb=ON_PRIMARY)
    add_rect(slide, R_x + 360000, R_y + 1278000, 360000, 18000, ON_PRIMARY)
    add_multiline(slide, [
        '• Headline — Pretendard 20pt, Bold, white',
        '• 소제목   — Pretendard 16pt, Bold, blue',
        '• 본문     — Pretendard 14pt, Regular, white',
        '• 카드 배경 — brand-teal #1a3a3a',
        '• 사용 — 핵심 결론 / Featured 솔루션',
    ], R_x + 360000, R_y + 1404000, R_w - 720000, R_h - 1764000,
       size=13, rgb=ON_PRIMARY)

    # Footer + page no
    add_text(slide, '장동인 | KAIST 김재철AI대학원 · AIBB LAB',
             864000, 6480000, 6048000, 270000, size=11, rgb=MUTED)
    add_text(slide, '03', 11556000, 162000, 432000, 270000,
             font='Inter', size=11, rgb=MUTED, align=PP_ALIGN.RIGHT)


def main():
    # 1. v1 원본을 template.pptx로 복원 (슬라이드 1, 2 콘텐츠 그대로)
    shutil.copy(V1_BACKUP, OUTPUT)
    print(f'복원: {V1_BACKUP} → {OUTPUT}')

    prs = Presentation(OUTPUT)
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 2. 모든 슬라이드 배경을 white로
    for i, slide in enumerate(prs.slides, start=1):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE
        print(f'슬라이드 {i:02d} 배경 white 적용')

    # 3. 슬라이드 3만 v2.0 콘텐츠 가이드로 재구성
    if len(prs.slides) >= 3:
        build_guide_slide(prs.slides[2])
        print('슬라이드 03 v2.0 콘텐츠 가이드 재구성 (삼각형 +90°)')

    prs.save(OUTPUT)
    print(f'저장 완료: {OUTPUT}')


if __name__ == '__main__':
    main()
