#!/usr/bin/env python3
"""Check that every concept is defined before it is first used in a deck.

usage:
  python scripts/check_terms.py deck.pptx terms.json [--notes] [--include-headings]
  python scripts/check_terms.py deck.pptx terms.json --discover   # 사전에 없는 영문 약어 후보 나열

terms.json:
  {"가상 팹": {"aliases": ["가상팹", "Virtual Fab"], "defined_in": "정의 슬라이드 제목의 일부"}, ...}

규칙: 용어(또는 별칭)가 처음 나오는 슬라이드가 그 용어를 정의하는 슬라이드보다 앞서면 실패한다.
제목(Title)이 없는 장(표지·프로필·Part 구분·Thank You)과 목차는 예고일 뿐이므로 기본 제외한다.
"""
import json
import re
import sys

from pptx import Presentation


def slide_texts(path, use_notes, include_headings):
    out = []
    for i, s in enumerate(Presentation(path).slides, 1):
        title = next((sh.text_frame.text for sh in s.shapes if sh.name == "Title" and sh.has_text_frame), "")
        if not include_headings and (not title or title.strip() in ("목차", "Contents", "Agenda")):
            continue
        t = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t.append(sh.text_frame.text)
            if sh.has_table:
                t += [c.text for r in sh.table.rows for c in r.cells]
        if use_notes and s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            t.append(s.notes_slide.notes_text_frame.text)
        out.append((i, title, "\n".join(t)))
    return out


def hit(text, name):
    if re.fullmatch(r"[A-Za-z0-9\-/ .]+", name):
        return re.search(r"(?<![A-Za-z0-9])" + re.escape(name) + r"(?![A-Za-z0-9])", text) is not None
    return name in text


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    deck, terms_path = sys.argv[1], sys.argv[2]
    slides = slide_texts(deck, "--notes" in sys.argv, "--include-headings" in sys.argv)
    terms = json.load(open(terms_path, encoding="utf-8"))
    if "--discover" in sys.argv:
        known = {n.lower() for k, v in terms.items() for n in [k, *v.get("aliases", [])]}
        seen = {}
        for i, _, tx in slides:
            for m in re.findall(r"(?<![A-Za-z0-9])[A-Z][A-Za-z0-9\-]*[A-Z0-9][A-Za-z0-9\-]*", tx):
                if m.lower() not in known and m not in seen:
                    seen[m] = i
        for k, v in seen.items():
            print(f"CANDIDATE  {k} (first {v})")
        print(f"{len(seen)} undefined acronym candidates; add real terms to {terms_path} or record why they are exempt")
        return 0
    bad = 0
    for term, spec in terms.items():
        names = [term, *spec.get("aliases", [])]
        defs = [i for i, title, _ in slides if spec["defined_in"] in title]
        uses = [i for i, _, tx in slides if any(hit(tx, n) for n in names)]
        if not defs:
            print(f'MISSING-DEF  {term}: no slide title contains "{spec["defined_in"]}"')
            bad += 1
            continue
        early = [u for u in uses if u < defs[0]]
        if early:
            print(f"USED-BEFORE  {term}: defined on slide {defs[0]}, used earlier on {early}")
            bad += 1
    print(f"{len(terms)} terms checked, {bad} problems")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
